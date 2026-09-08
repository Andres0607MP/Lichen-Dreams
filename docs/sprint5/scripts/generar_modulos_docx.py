# -*- coding: utf-8 -*-
"""Genera los 5 documentos DOCX del Sprint 5 — Lichen Dreams (por módulo).

Cada documento es autosuficiente y se basa en resultados REALES del Sprint 5,
con evidencias existentes en docs/sprint5/evidencias/.
Selenium NO forma parte de la validación oficial (solo nota de estrategia).
"""
import sys
from pathlib import Path
from datetime import date

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[3]               # raíz del proyecto
EVID = ROOT / "docs" / "sprint5" / "evidencias"
OUT_DIR = ROOT / "docs" / "sprint5"

AZUL = RGBColor(0x1F, 0x4E, 0x79)
VERDE = RGBColor(0x1E, 0x7B, 0x34)
GRIS = RGBColor(0x59, 0x59, 0x59)
FECHA = "5 de septiembre de 2026"

TECNICOS = {
    "autenticacion": {
        "archivo": "Sprint_5_Modulo_1_Autenticacion_Usuario.docx",
        "nombre": "Módulo 1 — Autenticación y gestión de usuario",
        "funcionalidades": [
            "Registro de cuenta (creación de usuario local)",
            "Inicio de sesión con correo y contraseña",
            "Inicio de sesión con Google (Google Sign-In)",
            "Cierre de sesión (revocación de sesiones)",
            "Tokens JWT y renovación de sesión (refresh)",
            "Recuperación de contraseña (código LCHN y código por correo)",
            "Cambio de contraseña y eliminación de cuenta",
            "Sesiones activas y su revocación",
            "Perfil del usuario (consulta, actualización, datos personales)",
            "Código de recuperación (generación/regeneración)",
            "Verificación de correo",
            "Administración de usuarios (funcionalidades de administración)",
        ],
        "alcance_incluye": [
            "Registro, inicio de sesión local y con Google, JWT/refresh, cierre de sesión.",
            "Recuperación de contraseña con código de recuperación LCHN (un solo uso) y con código temporal por correo.",
            "Perfil del usuario autenticado (consulta y actualización de datos).",
            "Cambio de contraseña, revocación de sesiones y eliminación de cuenta.",
            "Validaciones de formularios y de contrato API relacionadas con la cuenta.",
        ],
        "alcance_relaciona": [
            "La subida y el acceso a imágenes de perfil se apoyan en el servicio de imágenes (Módulo 2).",
            "El inicio de sesión habilita el acceso al dashboard y al historial (Módulos 3 y 5).",
            "Las funcionalidades de administración de usuarios se integran con los CRUD administrativos del backend.",
        ],
        "dependencias": ["Backend: rutas /auth y /profile", "FastAPI + JWT (python-jose) + Passlib/bcrypt", "Base de datos MySQL (usuarios, sesiones, recovery_codes) ", "Google Sign-In (cliente OAuth)"],
        "alcance_fuera": [
            "Validación visual del flujo completo de Google Sign-In en un dispositivo físico (requiere credenciales OAuth del dispositivo).",
            "Pruebas de carga o de seguridad avanzada más allá de la revisión automatizada realizada.",
        ],
        "tipos": ["Funcionales", "Autenticación/seguridad", "Negativas", "Casos límite", "Validación", "API", "Regresión", "Integración con backend"],
        "casos": [
            ["M1-01", "Registro exitoso", "No existir el correo", "Nombre, correo y contraseña válida (>=6 y carácter especial)", "1. Registrar usuario. 2. Verificar respuesta.", "201; código de recuperación LCHN"],
            ["M1-02", "Registro con correo inválido", "Formulario de registro", "Correo sin @", "1. Intentar registro. 2. Verificar respuesta.", "422 (validación de correo)"],
            ["M1-03", "Registro con contraseña débil", "Formulario de registro", "Contraseña menor a 6 caracteres", "1. Intentar registro. 2. Verificar respuesta.", "422 (contraseña mínima y carácter especial)"],
            ["M1-04", "Registro con correo duplicado", "Correo ya registrado", "Mismo correo", "1. Registrar o repetir correo existente.", "400 «Usuario ya existe»"],
            ["M1-05", "Registro con campos obligatorios vacíos", "Formulario de registro", "Petición sin datos", "1. Enviar petición vacía.", "422 (campos requeridos)"],
            ["M1-06", "Login exitoso", "Usuario registrado y activo", "Correo + contraseña correctos", "1. Enviar credenciales. 2. Verificar tokens.", "200; access_token y refresh_token"],
            ["M1-07", "Login con contraseña incorrecta", "Usuario existente", "Contraseña incorrecta", "1. Intentar login.", "401 «Credenciales inválidas»"],
            ["M1-08", "Login con usuario inexistente", "—", "Correo no registrado", "1. Intentar login.", "401 «Credenciales inválidas»"],
            ["M1-09", "Login con campos vacíos", "—", "Sin contraseña", "1. Intentar login sin password.", "422"],
            ["M1-10", "Sesión / autenticación válida", "Token emitido", "GET /auth/me", "1. Obtener token. 2. Consultar usuario actual.", "200; datos del usuario autenticado"],
            ["M1-11", "Token inválido rechazado", "—", "JWT inválido", "1. Enviar token inválido.", "401"],
            ["M1-12", "Solicitud de recuperación por correo", "Correo registrado", "POST /auth/forgot-password", "1. Solicitar recuperación.", "200; respuesta genérica (el envío SMTP real requiere SMTP configurado)"],
            ["M1-13", "Recuperación con código LCHN válido", "Código LCHN vigente", "POST /auth/recover-with-code", "1. Enviar código y nueva contraseña.", "200; permite login con nueva contraseña"],
            ["M1-14", "Código inválido en recuperación", "—", "Código incorrecto", "1. Resetear con código incorrecto.", "400 «Código inválido o expirado»"],
            ["M1-15", "Login con nueva contraseña", "Contraseña restablecida", "Nueva contraseña", "1. Iniciar sesión tras restablecer.", "200"],
            ["M1-16", "Perfil: carga de datos", "Sesión válida", "GET /profile", "1. Consultar perfil.", "200; datos del usuario"],
            ["M1-17", "Perfil: actualización de datos", "Sesión válida", "PUT /profile (nombre, teléfono)", "1. Enviar actualización.", "200; datos actualizados"],
            ["M1-18", "Perfil: correo inválido", "Sesión válida", "Correo mal formado", "1. Intentar actualizar correo.", "422"],
            ["M1-19", "Google Sign-In (contrato backend)", "Token simulado de Google", "POST /auth/google", "1. Ejecutar pruebas automatizadas de Google auth.", "Casos cubiertos por la suite automatizada (18 pruebas)"],
            ["M1-20", "Google Sign-In en dispositivo", "Dispositivo con credenciales OAuth", "Flujo de la pantalla de login (Google)", "1. Iniciar sesión con Google en dispositivo.", "No ejecutado / fuera del alcance del conjunto de pruebas ejecutado (requiere entorno OAuth de dispositivo)"],
            ["M1-21", "Cambio de contraseña y cierre de sesión", "Sesión activa", "POST /auth/change-password y /auth/logout", "1. Ejecutar pruebas automatizadas de sesiones/revocación.", "Cubierto por la suite backend"],
            ["M1-22", "Administración de usuarios (admin)", "Rol administrador", "Funcionalidades de administración de usuarios", "1. Evaluar endpoints de administración.", "No ejecutado / fuera del alcance del conjunto de pruebas ejecutado en la UI (endpoints disponibles)"],
        ],
        "matriz": [
            ["M1-01", "Registro exitoso", "201 + código LCHN", "201; recovery_code LCHN", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-02", "Registro correo inválido", "422", "422", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-03", "Registro contraseña débil", "422", "422", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-04", "Registro correo duplicado", "400", "400 «Usuario ya existe»", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-05", "Registro campos vacíos", "422", "422", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-06", "Login exitoso", "200 + tokens", "200; access_token", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-07", "Login contraseña incorrecta", "401", "401", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-08", "Login usuario inexistente", "401", "401", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-09", "Login campos vacíos", "422", "422", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-10", "Sesión válida (/auth/me)", "200", "200; datos del usuario", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-11", "Token inválido", "401", "401", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-12", "Solicitud recuperación por correo", "200 genérico", "200 genérico", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-13", "Recuperación con código LCHN", "200", "200; login posterior 200", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-14", "Código inválido", "400", "400", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-15", "Login con nueva contraseña", "200", "200", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-16", "Perfil: carga", "200 con datos", "200", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-17", "Perfil: actualización", "200 con cambios", "200; nombre actualizado", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-18", "Perfil: correo inválido", "422", "422", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M1-19", "Google Sign-In (backend)", "Casos de la suite Google", "18/18 pruebas de test_google_auth", "Aprobado", "backend/tests/test_google_auth.py"],
            ["M1-20", "Google Sign-In en dispositivo", "—", "—", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado", "Requiere credenciales OAuth de dispositivo"],
            ["M1-21", "Cambio de contraseña / cierre de sesión", "Cubierto por suite", "Aprobado por suite", "Aprobado", "backend/tests/test_auth_sessions.py, test_recovery_code.py"],
            ["M1-22", "Admin usuarios (UI)", "—", "—", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado", "Endpoints admin disponibles; UI administrador no evaluada en este conjunto"],
        ],
        "pytest": {
            "texto": "37 pruebas del backend están dedicadas específicamente a este módulo: test_auth_sessions (3), test_recovery_code (13), test_password (2), test_google_auth (18) y test_refresh_logout (1). El resto de la suite (incluidos test_endpoints, test_api y test_models) cubre la integración transversal. Suite completa: 176/176 aprobadas.",
        },
        "flutter": "Pruebas de widget relacionadas con la cuenta: auth_state_test, auth_isolation_test, google_login_state_test, google_sign_in_button_test, widget_test (pantalla de login) y recovery_code_screen_test. Total Flutter: 37/37 aprobadas; flutter analyze: 0 errores.",
        "e2e": "Dentro de la auditoría E2E vía API (38/38 aprobadas) corresponden a este módulo: registro, login, sesión, recuperación por código LCHN, perfil (carga/actualización/validación) y rechazo de tokens inválidos.",
        "crud": "Las funcionalidades de administración de usuarios (Módulo 1) están disponibles en el backend; su CRUD no fue ejercitado en el conjunto de pruebas ejecutado (se conserva como No ejecutado).",
        "bug": None,
        "evidencias": [
            "api/auditoria_sprint5_dllpsh.json (auditoría E2E definitiva, incluye registro/login/recuperación/perfil)",
            "api/ejecucion_auditoria_raw.txt (salida de consola de la auditoría)",
            "api/auditoria_sprint5_iu538v.json y api/auditoria_sprint5_pqqlqt.json (ejecuciones previas)",
        ],
        "conclusiones": [
            "Los flujos de autenticación (registro, login, sesión, recuperación, perfil) fueron validados con resultados reales vía API y mediante la suite automatizada del backend.",
            "Las validaciones de formulario y contrato (correo, contraseña, campos obligatorios) se comportan según lo esperado (422/400).",
            "El flujo de Google Sign-In está cubierto funcionalmente a nivel de backend por 18 pruebas automatizadas; la validación completa en dispositivo queda fuera del alcance del conjunto ejecutado.",
            "La recuperación por correo responde de forma genérica y segura (anti-enumeración); el envío SMTP real no se verificó en este entorno.",
        ],
        "integracion": "El módulo depende del backend FastAPI (/auth, /profile), de la persistencia en MySQL (usuarios, sesiones, recovery_codes) y del servicio JWT. Su correcto funcionamiento habilita el acceso al resto de los módulos del sistema.",
    },
    "captura_ia": {
        "archivo": "Sprint_5_Modulo_2_Captura_Analisis_IA.docx",
        "nombre": "Módulo 2 — Captura y análisis de imágenes con IA",
        "funcionalidades": [
            "Captura de imagen con la cámara del dispositivo",
            "Selección de imagen desde la galería",
            "Subida y validación de imágenes (extensión, MIME y magic bytes)",
            "Procesamiento del análisis con el modelo de IA",
            "Clasificación del líquen en 3 categorías (saludable, contaminado, desconocido)",
            "Manejo de imágenes que no corresponden a un líquen (derivan a «desconocido»)",
            "Control de errores de procesamiento (archivo inexistente/corrupto)",
            "Persistencia de análisis con su resultado, confianza e información ambiental",
            "Estado del análisis (completed / processing / error)",
            "Humedad, calidad del aire y recomendación asociadas",
            "Asociación de especie al análisis",
            "Servicio de imágenes (subida, consulta, eliminación) con control de acceso",
            "Gestión de especies (funcionalidades de administración relacionadas)",
        ],
        "alcance_incluye": [
            "Captura/selección de imagen (cámara/galería) en la aplicación.",
            "Subida, validación y procesamiento de la imagen contra el modelo de IA activo (v8).",
            "Contrato de análisis (resultado, confianza, estado, humedad, calidad del aire, recomendación).",
            "Persistencia de análisis y acceso a sus resultados.",
            "Servicio de imágenes y control de acceso (incluye la corrección BUG-001).",
        ],
        "alcance_relaciona": [
            "El análisis capturado se visualiza en Historial y Resultados (Módulo 3).",
            "La ubicación capturada se usa para asociar el análisis a un punto del mapa (Módulo 4).",
            "El catálogo de especies (Módulo 5) permite asociar una especie al análisis.",
        ],
        "dependencias": ["Modelo activo: lichen_model_v8.keras (registrado en la base de datos)", "TensorFlow/Keras + OpenCV", "Backend: rutas /analysis y /imagenes", "MySQL (analisis, procesamiento_ia, imagenes, notificaciones, historial_actividad)"],
        "alcance_fuera": [
            "Captura real mediante cámara en un dispositivo físico dentro de este conjunto de pruebas (no se automatizó la cámara).",
            "Entrenamiento o mejora del modelo de IA (no corresponde al Sprint 5).",
        ],
        "tipos": ["Funcionales", "IA / clasificación", "Negativas", "Casos límite", "Validación", "API", "Persistencia", "Regresión", "Seguridad (acceso a imágenes)"],
        "casos": [
            ["M2-01", "Subida de imagen válida", "Sesión autenticada", "Imagen JPEG", "1. Subir imagen. 2. Verificar respuesta.", "200; URL de imagen almacenada"],
            ["M2-02", "Análisis de líquen saludable", "Imagen del dataset (ls_0001)", "POST /analysis/process (gallery)", "1. Procesar imagen.", "200; resultado «liquen saludable»"],
            ["M2-03", "Análisis de imagen que no es un líquen", "Imagen sintética de prueba", "POST /analysis/process", "1. Procesar imagen.", "200; clasificada como «liquen desconocido» (modelo de 3 clases)"],
            ["M2-04", "Análisis sin imagen", "—", "Petición sin archivo", "1. Intentar procesar sin imagen.", "422 «Debes enviar una imagen o image_url»"],
            ["M2-05", "Análisis desde cámara sin ubicación", "image_source=camera sin GPS", "POST /analysis/process", "1. Intentar análisis de cámara sin ubicación.", "422; mensaje de GPS requerido"],
            ["M2-06", "Análisis persistido con ubicación", "Ubicación registrada", "POST /analysis/process (camera + id_ubicacion)", "1. Procesar con ubicación.", "200; id de análisis > 0"],
            ["M2-07", "Consulta de resultado persistido", "Análisis existente", "GET /analysis/results/{id}", "1. Consultar resultados.", "200; resultado y contrato completo"],
            ["M2-08", "Estado del análisis", "Análisis existente", "GET /analysis/{id}/status", "1. Consultar estado.", "200; status=completed"],
            ["M2-09", "Inferencias del modelo IA (3 clases + no-líquen)", "Imágenes reales del dataset", "predict() del clasificador", "1. Ejecutar inferencias reales. 2. Registrar confianza.", "Categorías y confianzas reales documentadas (9 inferencias)"],
            ["M2-10", "Manejo de errores del modelo", "Archivo inexistente / corrupto", "predict() sobre rutas inválidas", "1. Invocar con rutas inválidas.", "FileNotFoundError / ValueError controlados"],
            ["M2-11", "Control de acceso a imágenes (BUG-001)", "—", "GET /imagenes y DELETE /imagenes/{id}", "1. Probar sin token y con propiedad. 2. Verificar.", "401 sin token; 200/204 propietario; 403 ajeno"],
            ["M2-12", "Asociación de especie al análisis", "Análisis propio y especie registrada", "PUT /analysis/{id}/species", "1. Asociar especie. 2. Verificar.", "Cubierto por pruebas automatizadas (catálogo de especies)"],
            ["M2-13", "Captura con cámara en dispositivo", "Dispositivo Android con cámara", "Flujo de captura en pantalla de análisis", "1. Capturar imagen.", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado (requiere dispositivo)"],
            ["M2-14", "Selección desde galería en dispositivo", "Dispositivo Android", "Flujo de galería", "1. Seleccionar imagen.", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado (requiere dispositivo)"],
        ],
        "matriz": [
            ["M2-01", "Subida de imagen válida", "200", "200 (URL almacenada)", "Aprobado", "backend/tests (test_endpoints, test_api)"],
            ["M2-02", "Análisis líquen saludable", "200; saludable", "200; «liquen saludable», confianza 0.9997", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M2-03", "Análisis imagen no-líquen", "200; clasificación 3 clases", "200; «liquen desconocido», confianza 1.0", "Aprobado", "api/auditoria_sprint5_*.json, ia/resultados_ia.json"],
            ["M2-04", "Análisis sin imagen", "422", "422", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M2-05", "Cámara sin ubicación", "422 GPS", "422 (GPS requerido)", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M2-06", "Análisis persistido con ubicación", "id > 0", "id=31/32; saludable", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M2-07", "Consulta de resultado persistido", "200", "200", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M2-08", "Estado del análisis", "completed", "status=completed", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M2-09", "Inferencias del modelo IA", "9 inferencias reales", "9 OK (saludable/contaminado/desconocido/no-líquen)", "Aprobado", "ia/resultados_ia.json, ia/ejecucion_ia.txt"],
            ["M2-10", "Errores controlados del modelo", "Errores esperados", "FileNotFoundError / ValueError controlados", "Aprobado", "ia/resultados_ia.json"],
            ["M2-11", "Control de acceso a imágenes (BUG-001)", "401/200/204/403", "401 sin token; 200/204 propietario; 403 ajeno", "Aprobado", "correcciones/resultado_correccion_bug001.json"],
            ["M2-12", "Asociación de especie", "Cubierto por suite", "Pruebas de catálogo de especies (asignación)", "Aprobado", "backend/tests/test_catalog_zones.py"],
            ["M2-13", "Cámara en dispositivo", "—", "—", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado", "Requiere dispositivo Android con cámara"],
            ["M2-14", "Galería en dispositivo", "—", "—", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado", "Requiere dispositivo Android"],
        ],
        "pytest": {
            "texto": "81 pruebas del backend están dedicadas específicamente a este módulo: test_analisis (4), test_analisis_desconocido (9), test_lichen_classifier (6), test_inferencia_v3 (14), test_issue14_mock_integration (1), test_image_access_control (6), test_v7_split_pipeline (16), test_revisor_94 (10), test_analysis_contract (8) y test_species_schemas (7). Suite completa: 176/176 aprobadas.",
        },
        "flutter": "Pruebas de widget relacionadas con el análisis y la gestión de especies: species_form_dialog_test. Total Flutter: 37/37 aprobadas; flutter analyze: 0 errores.",
        "e2e": "Dentro de la auditoría E2E vía API (38/38 aprobadas) corresponden a este módulo: análisis de imagen saludable, análisis de imagen no-líquen, análisis sin imagen (422), análisis de cámara sin ubicación (422), análisis persistido con ubicación, consulta de resultados y estado del análisis.",
        "crud": "La gestión administrativa de especies (crear, consultar, actualizar, eliminar y duplicado 409) formó parte de la auditoría CRUD administrativa (21/21 aprobadas); el catálogo de especies es consumido por el flujo de análisis.",
        "bug": {
            "id": "BUG-001",
            "modulo": "Imágenes",
            "descripcion": "Los endpoints /imagenes, /imagenes/{id} y DELETE /imagenes/{id} no exigían autenticación y exponían rutas de imágenes privadas (analyses/user_N), permitiendo además su eliminación sin sesión.",
            "severidad": "Alta",
            "prioridad": "Alta",
            "correccion": "Se protegió GET /imagenes, GET /imagenes/{id} y DELETE /imagenes/{id} con autenticación obligatoria y control de propiedad/permiso (CAN_VIEW_PRIVATE_IMAGES).",
            "revalidacion": "Tras la corrección: sin token: 401; propietario: 200/204; otro usuario (admin sin permiso especial): 403. Se agregó prueba de regresión y la suite pasó 176/176.",
        },
        "evidencias": [
            "ia/resultados_ia.json (inferencias reales del modelo activo)",
            "ia/ejecucion_ia.txt (salida de consola de las inferencias)",
            "ia/imagen_no_liquen.png (imagen sintética de prueba utilizada en la inferencia no-líquen)",
            "api/auditoria_sprint5_dllpsh.json (análisis E2E)",
            "correcciones/resultado_correccion_bug001.json y correcciones/ejecucion_bug001.txt (validación BUG-001)",
        ],
        "imagen_ia": True,
        "conclusiones": [
            "El pipeline de análisis con IA se validó con resultados reales: el modelo activo (v8) es realmente el registrado en la base de datos y se ejecutaron 9 inferencias reales con 2 errores controlados.",
            "Las imágenes de líquen saludable se clasificaron correctamente; una imagen de la clase «contaminado» derivó a «desconocido» (comportamiento real documentado) y las imágenes no-líquen se clasificaron como «liquen desconocido», sin convertirse automáticamente en «no es un líquen».",
            "La persistencia del análisis (id real, historial, resultados consultables) fue confirmada.",
            "Se corrigió y revalidó BUG-001 (control de acceso de imágenes); la suite pasó 176/176 y la corrección quedó cubierta por una prueba de regresión.",
            "La captura mediante cámara y la selección desde galería en un dispositivo físico quedaron documentadas como fuera del alcance de este conjunto de pruebas (requieren entorno de dispositivo).",
        ],
        "integracion": "El módulo se integra con la IA (modelo v8 + OpenCV), con FastAPI (/analysis, /imagenes), con MySQL (registros de análisis, procesamiento e imágenes) y alimenta el historial, el dashboard, el mapa y las notificaciones del sistema.",
    },
    "resultados_historial": {
        "archivo": "Sprint_5_Modulo_3_Resultados_Historial_Reportes.docx",
        "nombre": "Módulo 3 — Resultados, historial y reportes",
        "funcionalidades": [
            "Visualización de resultados completos de un análisis",
            "Contrato de respuesta del análisis (resultado, confianza, estado, humedad, calidad del aire, recomendación)",
            "Historial de análisis del usuario",
            "Historial vacío / sin análisis",
            "Eliminación de análisis e historial",
            "Dashboard de estadísticas sincronizadas con historial",
            "Reportes ambientales (creación, listado, detalle, eliminación)",
            "Persistencia y consulta de resultados asociados al usuario",
            "Notificaciones de análisis completado",
        ],
        "alcance_incluye": [
            "Consulta de resultados y estado de análisis desde la API y el historial.",
            "Listado de historial propio y consulta de análisis del usuario.",
            "Estadísticas del dashboard alimentadas por el historial/análisis.",
            "Reportes ambientales generados a partir de las estadísticas del usuario.",
            "Eliminación de análisis con su historial asociado.",
        ],
        "alcance_relaciona": [
            "Los resultados provienen del módulo de análisis (Módulo 2).",
            "El dashboard también muestra indicadores derivados de ubicaciones (Módulo 4).",
            "Las notificaciones de análisis completado se generan durante el procesamiento (Módulo 2).",
        ],
        "dependencias": ["Backend: rutas /history, /analysis (consulta), /dashboard y /reports", "MySQL (historial_actividad, analisis, reportes)", "FastAPI + SQLAlchemy"],
        "alcance_fuera": [
            "Visualización final en pantalla dentro de un dispositivo físico (validación de la interfaz del historial/reportes no incluida en este conjunto de pruebas).",
            "Generación de reportes visuales o exportaciones (no implementadas en el alcance evaluado).",
        ],
        "tipos": ["Funcionales", "Negativas", "Casos límite", "Persistencia", "API", "Integración", "Regresión"],
        "casos": [
            ["M3-01", "Historial con análisis", "Existen análisis del usuario", "GET /history", "1. Consultar historial.", "200; lista con al menos un análisis"],
            ["M3-02", "Historial vacío", "Usuario sin análisis", "GET /history", "1. Consultar historial de usuario nuevo.", "200; lista vacía"],
            ["M3-03", "Dashboard sincronizado", "Análisis previo", "GET /dashboard/stats", "1. Consultar estadísticas.", "200; contadores coherentes con el historial"],
            ["M3-04", "Consulta de resultados por id", "Análisis existente", "GET /analysis/results/{id}", "1. Consultar resultados.", "200; contrato completo"],
            ["M3-05", "Eliminación de análisis", "Análisis propio y ajeno", "DELETE /analysis/{id}", "1. Eliminar análisis propio. 2. Intentar ajeno.", "204 propio; 403 ajeno"],
            ["M3-06", "Reporte ambiental: creación", "Usuario con análisis", "POST /reports/environmental", "1. Generar reporte.", "Creado con datos de estadísticas"],
            ["M3-07", "Reporte: listado", "Reportes existentes", "GET /reports", "1. Consultar reportes.", "200"],
            ["M3-08", "Reporte: eliminación", "Reporte propio/ajeno", "DELETE /reports/{id}", "1. Eliminar propio. 2. Intentar ajeno.", "204 propio; 404 ajeno"],
            ["M3-09", "Reporte sin análisis", "Usuario sin análisis", "POST /reports/environmental", "1. Intentar generar reporte.", "400"],
            ["M3-10", "Notificación de análisis completado", "Análisis persistido", "GET /notificaciones", "1. Consultar notificaciones.", "200; incluye «Tu análisis está listo»"],
            ["M3-11", "Consulta de análisis propios", "Sesión válida", "GET /analysis/my", "1. Consultar análisis propios.", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado"],
        ],
        "matriz": [
            ["M3-01", "Historial con análisis", "200; items=1", "200; items=1", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M3-02", "Historial vacío", "200; items=0", "200; items=0 (usuario nuevo)", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M3-03", "Dashboard sincronizado", "200 con contadores", "200; analysis_count=1, healthy_count=1", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M3-04", "Resultados por id", "200", "200; contrato completo", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M3-05", "Eliminación de análisis", "204 propio / 403 ajeno", "Cubierto por suite backend", "Aprobado", "backend/tests/test_endpoints.py"],
            ["M3-06", "Reporte ambiental", "Creación con datos", "Cubierto por suite backend", "Aprobado", "backend/tests/test_reports.py"],
            ["M3-07", "Reporte: listado", "200", "Cubierto por suite backend", "Aprobado", "backend/tests/test_reports.py"],
            ["M3-08", "Reporte: eliminación", "204/404", "Cubierto por suite backend", "Aprobado", "backend/tests/test_reports.py"],
            ["M3-09", "Reporte sin análisis", "400", "Cubierto por suite backend", "Aprobado", "backend/tests/test_reports.py"],
            ["M3-10", "Notificación análisis listo", "200", "200; «Tu análisis está listo»", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M3-11", "Consulta análisis propios (UI)", "—", "—", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado", "Endpoint disponible; no invocado en el conjunto ejecutado"],
        ],
        "pytest": {
            "texto": "Pruebas del backend específicas de este módulo: test_dashboard_stats_sync (2) y test_reports (6); los casos de eliminación con control de propiedad y de contrato de historial/respuesta se cubren dentro de test_endpoints (23) y test_analysis_contract (8, compartido con el Módulo 2). Suite completa: 176/176 aprobadas.",
        },
        "flutter": "Pruebas de widget relacionadas con el dashboard y la sincronización historial/dashboard: stats_sync_test y dashboard_semantics_regression_test. Total Flutter: 37/37 aprobadas; flutter analyze: 0 errores.",
        "e2e": "Dentro de la auditoría E2E vía API (38/38 aprobadas) corresponden a este módulo: carga de historial, estadísticas del dashboard y notificaciones de análisis completado.",
        "crud": "La generación, consulta y eliminación de reportes ambientales fue validada por la suite de pruebas del backend (6 pruebas aprobadas).",
        "bug": None,
        "evidencias": [
            "api/auditoria_sprint5_dllpsh.json (historial, dashboard, notificaciones)",
            "api/ejecucion_auditoria_raw.txt (salida de consola)",
            "backend/tests/test_reports.py y test_dashboard_stats_sync.py (pruebas ejecutadas)",
        ],
        "conclusiones": [
            "El historial y el dashboard se comportan de forma coherente: los contadores se derivan del historial/análisis real del usuario.",
            "Los reportes ambientales fueron validados por pruebas automatizadas (creación, listado, detalle y eliminación con control de propiedad).",
            "La notificación de análisis completado se genera y se consulta correctamente.",
            "La consulta «mis análisis» en la interfaz no formó parte del conjunto de pruebas ejecutado y queda documentada como fuera del alcance.",
        ],
        "integracion": "Este módulo consume los resultados generados por el Módulo 2 (análisis/IA), mantiene el historial en MySQL y alimenta el dashboard; los reportes dependen de los datos analíticos del usuario, cerrándose el ciclo Flutter → FastAPI → MySQL → IA.",
    },
    "ubicacion_mapa": {
        "archivo": "Sprint_5_Modulo_4_Ubicacion_Mapa.docx",
        "nombre": "Módulo 4 — Ubicación y mapa",
        "funcionalidades": [
            "Obtención de ubicación (GPS/geolocalización)",
            "Geocodificación inversa (dirección a partir de coordenadas)",
            "Registro y reutilización de ubicaciones (find-or-create)",
            "Persistencia del análisis asociado a una ubicación",
            "Visualización de puntos en el mapa (propios y comunitarios)",
            "Círculos de radio (5 m) alrededor de los puntos",
            "Filtros por nivel de calidad en el mapa",
            "Zonas ambientales (gestión administrativa e indicadores)",
            "Catálogo de zonas ambientales",
        ],
        "alcance_incluye": [
            "Registro de ubicación mediante find-or-create con radio de coincidencia.",
            "Publicación de puntos del mapa a partir de análisis con ubicación.",
            "Indicadores ambientales calculados a partir de los análisis dentro del radio.",
            "Funcionalidades administrativas de zonas ambientales.",
        ],
        "alcance_relaciona": [
            "Los análisis se capturan en el Módulo 2 y requieren ubicación cuando se usan desde cámara.",
            "Los puntos del mapa dependen del historial/análisis (Módulo 3).",
            "Los ajustes de privacidad de análisis compartidos se gestionan junto con la comunidad (Módulo 5).",
        ],
        "dependencias": ["Google Maps (clave configurada)", "Geolocator / Geocoding (frontend)", "Backend: /location, /api/maps, /admin/zones, /catalog/zones", "MySQL (ubicaciones, analisis_zonas_ambientales)"],
        "alcance_fuera": [
            "Carga del mapa y obtención de GPS en un dispositivo físico dentro del conjunto de pruebas ejecutado (requiere emulador/dispositivo con Google Maps).",
            "Escenarios de mapa con múltiples análisis en la misma coordenada (no ejecutados en el conjunto).",
        ],
        "tipos": ["Funcionales", "Negativas", "Casos límite", "Persistencia", "API", "Integración", "Regresión"],
        "casos": [
            ["M4-01", "Registro y reutilización de ubicación", "—", "POST /location/find-or-create (lat, lng)", "1. Registrar ubicación. 2. Verificar.", "200; id_ubicacion y bandera existed"],
            ["M4-02", "Persistencia de análisis con ubicación", "Ubicación creada", "POST /analysis/process (camera + id_ubicacion)", "1. Procesar con ubicación.", "200; id real persistido"],
            ["M4-03", "Puntos del mapa", "Análisis con ubicación", "GET /api/maps/points", "1. Consultar puntos.", "200; puntos con lat/lng y calidad"],
            ["M4-04", "Visibilidad propios/comunidad", "Análisis privado", "Puntos del mapa", "1. Consultar puntos propios y comunitarios.", "Propios visibles; comunidad sin datos en el conjunto ejecutado"],
            ["M4-05", "Zonas ambientales: creación (admin)", "Rol administrador", "POST /admin/zones", "1. Crear zona.", "201; indicadores calculados"],
            ["M4-06", "Zonas: edición (admin)", "Zona existente", "PUT /admin/zones/{id}", "1. Editar zona.", "200"],
            ["M4-07", "Zonas: eliminación (admin)", "Zona existente", "DELETE /admin/zones/{id}", "1. Eliminar zona.", "204"],
            ["M4-08", "Catálogo de zonas", "Zona creada", "GET /catalog/zones", "1. Consultar zonas.", "200; contiene zona con indicadores"],
            ["M4-09", "Círculos de radio y selección de puntos (UI)", "Lógica del frontend", "Tests de widget del mapa", "1. Ejecutar pruebas de widget.", "Lógica validada (10 pruebas de widget)"],
            ["M4-10", "Dos análisis en la misma coordenada", "2 análisis en coords iguales", "—", "1. Crear dos análisis en el mismo punto.", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado"],
            ["M4-11", "Carga del mapa en dispositivo", "Dispositivo/emulador con Google Maps", "Pantalla de mapa", "1. Abrir mapa.", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado"],
        ],
        "matriz": [
            ["M4-01", "Find-or-create ubicación", "200; id_ubicacion", "200; id_ubicacion=3, existed=false", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M4-02", "Análisis persistido con ubicación", "id>0", "id=31/32; saludable + humedad", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M4-03", "Puntos del mapa", "200 con puntos", "200; 1 punto con lat/lng y calidad", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M4-04", "Visibilidad propios/comunidad", "Propios visibles", "1 punto propio; comunidad sin datos", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M4-05", "Crear zona (admin)", "201; indicadores", "201; total_analisis=1, saludables=1", "Aprobado", "api/auditoria_crud_admin.json"],
            ["M4-06", "Editar zona (admin)", "200", "200", "Aprobado", "api/auditoria_crud_admin.json"],
            ["M4-07", "Eliminar zona (admin)", "204", "204", "Aprobado", "api/auditoria_crud_admin.json"],
            ["M4-08", "Catálogo de zonas", "200", "200; contiene la zona", "Aprobado", "api/auditoria_crud_admin.json"],
            ["M4-09", "Círculos y selección (widget)", "Pruebas de widget", "10 pruebas de widget del mapa", "Aprobado", "frontend/test/map_circles_test.dart, map_screen_selection_test.dart"],
            ["M4-10", "Dos análisis misma coordenada", "—", "—", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado", "No se crearon dos análisis coincidentes en el conjunto ejecutado"],
            ["M4-11", "Mapa en dispositivo", "—", "—", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado", "Requiere emulador/dispositivo con Google Maps"],
        ],
        "pytest": {
            "texto": "Pruebas del backend específicas de este módulo: test_catalog_zones (5) y test_zone_integration (5). La integración zona/ubicación también se cubre en test_endpoints (23, compartido). Suite completa: 176/176 aprobadas.",
        },
        "flutter": "Pruebas de widget del mapa: map_circles_test (círculos de radio y colores), map_circles_from_backend_test (parseo de puntos) y map_screen_selection_test (selección de puntos). Total Flutter: 37/37 aprobadas; flutter analyze: 0 errores.",
        "e2e": "Dentro de la auditoría E2E vía API (38/38 aprobadas) corresponden a este módulo: creación y reutilización de ubicaciones y obtención de puntos del mapa con sus análisis.",
        "crud": "La gestión administrativa de zonas ambientales (crear, editar, eliminar y catálogo) formó parte de la auditoría CRUD administrativa (21/21 aprobadas).",
        "bug": None,
        "evidencias": [
            "api/auditoria_sprint5_dllpsh.json (ubicación y puntos del mapa)",
            "api/auditoria_crud_admin.json y api/ejecucion_crud_admin.txt (zonas ambientales)",
            "frontend/test/map_circles_test.dart, map_circles_from_backend_test.dart y map_screen_selection_test.dart",
        ],
        "conclusiones": [
            "El registro y la reutilización de ubicaciones funciona correctamente (find-or-create con radio).",
            "Los puntos del mapa se generan a partir de análisis persistidos con ubicación, incluyendo su calidad.",
            "Las zonas ambientales se administran correctamente y sus indicadores se calculan a partir de los análisis reales.",
            "Los escenarios dependientes de GPS/Google Maps en dispositivo y los puntos coincidentes quedan documentados como fuera del alcance del conjunto ejecutado.",
        ],
        "integracion": "El módulo depende de la geolocalización del dispositivo y de Google Maps, y se integra con el análisis (Módulo 2) y el historial (Módulo 3) para poblar los puntos del mapa y los indicadores de zonas.",
    },
    "liquenpedia_comunidad": {
        "archivo": "Sprint_5_Modulo_5_LiquenPedia_Comunidad_Complementarios.docx",
        "nombre": "Módulo 5 — LiquenPedia, comunidad y funcionalidades complementarias",
        "funcionalidades": [
            "LiquenPedia: listado público de artículos",
            "LiquenPedia: detalle de artículo e información del autor",
            "LiquenPedia: búsqueda",
            "Categorías de artículos",
            "Creación, edición y eliminación de artículos (administración)",
            "Estados de publicación (draft/published/archived)",
            "Notificaciones del sistema",
            "Funcionalidades de comunidad (visibilidad y análisis compartidos)",
            "Catálogo de especies (referencia)",
            "Ajustes de la aplicación (tema, sonido, información, licencias, ayuda)",
        ],
        "alcance_incluye": [
            "Listado, búsqueda, detalle y CRUD administrativo de artículos LiquenPedia.",
            "Categorías de artículos.",
            "Notificaciones del sistema (listado y envío administrativo).",
            "Catálogo de especies como referencia.",
            "Ajustes de la aplicación y pantallas complementarias (licencias, ayuda, términos).",
            "Privacidad de análisis (compartidos) cuando aplica.",
        ],
        "alcance_relaciona": [
            "Las notificaciones de análisis completado provienen del Módulo 2.",
            "Los análisis compartidos se visualizan en el mapa (Módulo 4).",
            "El catálogo de especies se consume en el flujo de análisis (Módulo 2).",
        ],
        "dependencias": ["Backend: /liquenpedia, /categorias-liquenpedia, /catalog/species, /notificaciones y /admin/notifications", "MySQL (liquenpedia, categorias_articulos, notificaciones)", "FastAPI"],
        "alcance_fuera": [
            "Consumo de artículos en la interfaz móvil dentro de un dispositivo físico (validación de la UI no incluida en este conjunto).",
            "Funcionalidades sociales no implementadas (por ejemplo, comentarios o perfiles públicos).",
        ],
        "tipos": ["Funcionales", "Negativas", "Casos límite", "Persistencia", "API", "Integración", "Regresión"],
        "casos": [
            ["M5-01", "LiquenPedia: listado público", "Artículo publicado", "GET /liquenpedia", "1. Consultar listado.", "200; incluye el artículo publicado"],
            ["M5-02", "LiquenPedia: detalle", "Artículo existente", "GET /liquenpedia/{id}", "1. Consultar detalle.", "200; contenido y autor"],
            ["M5-03", "LiquenPedia: búsqueda", "Artículos existentes", "GET /liquenpedia?busqueda=...", "1. Buscar.", "200; resultados coincidentes"],
            ["M5-04", "Estado draft no visible", "Artículo en draft", "Listado público y detalle", "1. Consultar sin autenticación.", "No visible; 403 en detalle"],
            ["M5-05", "CRUD de artículos (admin)", "Rol administrador", "POST/PUT/DELETE /liquenpedia", "1. Crear/editar/eliminar.", "201/200/204"],
            ["M5-06", "Categorías de artículos", "Categorías registradas", "GET /categorias-liquenpedia", "1. Consultar categorías.", "200"],
            ["M5-07", "Envío de notificación (admin)", "Rol administrador", "POST /admin/notifications (all)", "1. Enviar a todos.", "201; count > 0"],
            ["M5-08", "Notificaciones del usuario", "Usuario con notificaciones", "GET /notificaciones", "1. Consultar.", "200"],
            ["M5-09", "Marcar leída / limpiar notificaciones", "Notificaciones existentes", "PATCH /notificaciones/{id}/read y DELETE /notificaciones/clear", "1. Marcar/limpiar.", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado"],
            ["M5-10", "Catálogo de especies", "Especies registradas", "GET /catalog/species", "1. Consultar catálogo.", "200; incluye especies creadas"],
            ["M5-11", "Ajustes: pantalla de licencias", "Widget de ajustes", "Pantalla de licencias", "1. Ejecutar prueba de widget.", "Renderizado correcto (test de widget)"],
            ["M5-12", "Ajustes en dispositivo (tema/sonido/ayuda)", "Dispositivo Android", "Pantallas de ajustes", "1. Navegar en dispositivo.", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado"],
        ],
        "matriz": [
            ["M5-01", "Listado público de artículos", "200; visible", "200; total=1 tras crear artículo", "Aprobado", "api/auditoria_crud_admin.json"],
            ["M5-02", "Detalle de artículo", "200", "200", "Aprobado", "api/auditoria_crud_admin.json"],
            ["M5-03", "Búsqueda", "200", "200", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M5-04", "Draft oculto", "No visible; 403", "visible=False; 403", "Aprobado", "api/auditoria_crud_admin.json"],
            ["M5-05", "CRUD artículos (admin)", "201/200/204", "201; 200; 204; 404 tras eliminar", "Aprobado", "api/auditoria_crud_admin.json"],
            ["M5-06", "Categorías de artículos", "200", "200 (categorías existentes)", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M5-07", "Notificación a todos (admin)", "201; count>0", "201; count=6", "Aprobado", "api/auditoria_crud_admin.json"],
            ["M5-08", "Notificaciones del usuario", "200", "200", "Aprobado", "api/auditoria_sprint5_*.json"],
            ["M5-09", "Marcar/limpiar notificaciones", "—", "—", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado", "Endpoints disponibles; no invocados en el conjunto ejecutado"],
            ["M5-10", "Catálogo de especies", "200", "200; contiene especies creadas", "Aprobado", "api/auditoria_crud_admin.json"],
            ["M5-11", "Pantalla de licencias", "Render sin errores", "Test de widget aprobado", "Aprobado", "frontend/test/licenses_screen_test.dart"],
            ["M5-12", "Ajustes en dispositivo", "—", "—", "No ejecutado / Fuera del alcance del conjunto de pruebas ejecutado", "Requiere dispositivo Android"],
        ],
        "pytest": {
            "texto": "Las funcionalidades de LiquenPedia y categorías se cubren transversalmente dentro de test_endpoints (23) y test_api (7), incluyendo la creación de artículos y el respeto del estado de publicación. Suite completa: 176/176 aprobadas.",
        },
        "flutter": "Prueba de widget relacionada con las pantallas complementarias: licenses_screen_test. Total Flutter: 37/37 aprobadas; flutter analyze: 0 errores.",
        "e2e": "Dentro de la auditoría E2E vía API (38/38 aprobadas) corresponden a este módulo: listado y búsqueda de LiquenPedia y categorías de artículos.",
        "crud": "La auditoría CRUD administrativa (21/21 aprobadas) cubrió en este módulo: creación/edición/eliminación de artículos (incluido draft), gestión de especies (catálogo) y envío de notificaciones a todos los usuarios.",
        "bug": None,
        "evidencias": [
            "api/auditoria_crud_admin.json y api/ejecucion_crud_admin.txt (artículos, notificaciones, especies)",
            "api/auditoria_sprint5_dllpsh.json (liquenpedia, categorías, notificaciones)",
            "frontend/test/licenses_screen_test.dart",
        ],
        "conclusiones": [
            "LiquenPedia permite listar, buscar y consultar artículos; el CRUD administrativo funciona correctamente y el estado draft se respeta (oculto al público, 403 en detalle).",
            "Las notificaciones se generan y consultan correctamente; el envío administrativo a todos los usuarios funcionó.",
            "El catálogo de especies está disponible y alimenta el flujo de análisis.",
            "Las funcionalidades dependientes de la interfaz móvil de ajustes y la gestión de leer/limpiar notificaciones quedaron documentadas como fuera del alcance del conjunto ejecutado.",
        ],
        "integracion": "Este módulo complementa el sistema: LiquenPedia y catálogos se consumen desde el dashboard y el análisis, las notificaciones acompañan el ciclo de análisis y los ajustes completan la experiencia del usuario móvil.",
    },
}


def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def add_header_footer(doc):
    section = doc.sections[0]
    hdr = section.header.paragraphs[0]
    hdr.text = "Lichen Dreams — Sprint 5 · Pruebas y Estabilización del Software"
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hdr.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = GRIS
        r.font.italic = True
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Página ")
    run.font.size = Pt(8)
    run.font.color.rgb = GRIS
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r1 = p.add_run(); r1.font.size = Pt(8); r1._r.append(f1)
    r2 = p.add_run(); r2._r.append(it)
    r3 = p.add_run(); r3.font.size = Pt(8); r3._r.append(f2)
    for rr in [r1, r2, r3]:
        rr.font.color.rgb = GRIS


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = AZUL
    return h


def table_(doc, headers, rows, widths=None, font_size=8.0, header_fill="1F4E79"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htext in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(htext)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_bg(cell, header_fill)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            run = cells[j].paragraphs[0].add_run(str(val))
            run.font.size = Pt(font_size)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Inches(w)
    return t


def figure(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    try:
        run.add_picture(str(path), width=Inches(5.2))
    except Exception as e:
        run.add_text(f"[Imagen no disponible: {path.name} — {e}]")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(caption)
    cr.font.size = Pt(8)
    cr.font.italic = True
    cr.font.color.rgb = GRIS


def build_module_doc(mod_key):
    cfg = TECNICOS[mod_key]
    doc = Document()
    add_header_footer(doc)

    # ---------- Portada ----------
    doc.add_paragraph(); doc.add_paragraph()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("LICHEN DREAMS"); run.font.size = Pt(34); run.bold = True; run.font.color.rgb = AZUL
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("«Lee el aire, entiende tu entorno»"); sr.font.size = Pt(15); sr.italic = True; sr.font.color.rgb = VERDE
    doc.add_paragraph()
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(cfg["nombre"]); tr.font.size = Pt(20); tr.bold = True
    doc.add_paragraph()
    pt = doc.add_paragraph(); pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pr = pt.add_run("Plan de pruebas — Sprint 5 (Pruebas y Estabilización del Software)"); pr.font.size = Pt(13); pr.font.color.rgb = GRIS
    doc.add_paragraph()
    doc.add_paragraph()
    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        f"Tecnología principal: Flutter\nPlataforma objetivo: Android\n\n"
        f"Aprendices SENA · Análisis y Desarrollo de Software (ADSO)\n"
        "Equipo: Hugo Andres Mancera Perez, Daniel Camilo Luque Briceño, Saira Yineth Aragon Suarez, "
        "Neyireth Dayana Soriano Ruiz, Heidy Lizeth Vivas Ramirez\n\n"
        f"Fecha: {FECHA}"
    )
    mr.font.size = Pt(10)
    mr.font.color.rgb = GRIS
    doc.add_page_break()

    # ---------- 1. Información general ----------
    heading(doc, "1. Información general del módulo", 1)
    info = [
        ["Proyecto", "Lichen Dreams"],
        ["Tipo de sistema", "Aplicación móvil (Android)"],
        ["Tecnología principal", "Flutter"],
        ["Backend", "FastAPI"],
        ["Base de datos", "MySQL"],
        ["Autenticación", "JWT / Google Sign-In"],
        ["IA", "Modelo para análisis de imágenes de líquenes"],
        ["API", "REST"],
        ["Módulo", cfg["nombre"]],
        ["Plataforma objetivo", "Android"],
    ]
    table_(doc, ["Campo", "Información"], info, widths=[2.2, 4.6])
    doc.add_paragraph()
    doc.add_paragraph(
        "Flutter es un framework multiplataforma; sin embargo, Lichen Dreams se desarrolla y "
        "valida como producto móvil Android en este proyecto. El soporte de otras plataformas no "
        "forma parte del alcance funcional ni de validación de este proyecto."
    )

    # ---------- 2. Objetivo ----------
    heading(doc, "2. Objetivo del módulo", 1)
    doc.add_paragraph(
        "Validar que las funcionalidades del módulo «" + cfg["nombre"] + "» se comporten de forma "
        "estable y conforme a lo esperado, que se integren correctamente con el backend (FastAPI) y "
        "la base de datos (MySQL), y que sus resultados queden correctamente persistidos, dentro "
        "del contexto de estabilización del Sprint 5."
    )

    # ---------- 3. Alcance ----------
    heading(doc, "3. Alcance del módulo", 1)
    heading(doc, "3.1 Funcionalidades incluidas", 2)
    for x in cfg["alcance_incluye"]:
        doc.add_paragraph(x, style="List Bullet")
    heading(doc, "3.2 Funcionalidades relacionadas y dependencias", 2)
    for x in cfg["alcance_relaciona"]:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_paragraph("Dependencias:" , style="List Bullet")
    for x in cfg["dependencias"]:
        doc.add_paragraph(x, style="List Bullet 2")
    heading(doc, "3.3 Fuera del alcance del módulo", 2)
    for x in cfg["alcance_fuera"]:
        doc.add_paragraph(x, style="List Bullet")

    # ---------- 4. Funcionalidades a probar ----------
    heading(doc, "4. Funcionalidades a probar", 1)
    rows = [[f"F{mod_key.upper()}-{i+1:02d}", f] for i, f in enumerate(cfg["funcionalidades"])]
    table_(doc, ["ID", "Funcionalidad"], rows, widths=[1.4, 5.4], font_size=8.0)

    # ---------- 5. Tipos de prueba ----------
    heading(doc, "5. Tipos de prueba del módulo", 1)
    for x in cfg["tipos"]:
        doc.add_paragraph(x, style="List Bullet")

    # ---------- 6. Diseño de casos ----------
    heading(doc, "6. Diseño de casos de prueba", 1)
    table_(doc, ["ID", "Caso de prueba", "Precondición", "Datos de entrada", "Pasos", "Resultado esperado"],
           cfg["casos"], widths=[0.7, 1.35, 1.1, 1.2, 1.1, 1.35], font_size=6.5)

    # ---------- 7. Matriz ----------
    heading(doc, "7. Matriz de ejecución de casos", 1)
    doc.add_paragraph(
        "Estados permitidos: Aprobado / Fallido / Bloqueado. Los casos sin evidencia de ejecución "
        "en el conjunto de pruebas del Sprint 5 se indican como «No ejecutado / Fuera del alcance "
        "del conjunto de pruebas ejecutado»."
    )
    table_(doc, ["ID Caso", "Caso ejecutado", "Resultado esperado", "Resultado obtenido", "Estado", "Evidencia"],
           cfg["matriz"], widths=[0.8, 1.5, 1.2, 1.4, 1.2, 1.3], font_size=6.5)

    # ---------- 8. Resultados automatizados ----------
    heading(doc, "8. Resultados automatizados existentes", 1)
    res = [
        ["Backend (pytest)", "176/176 pruebas aprobadas"],
        ["Flutter (flutter test)", "37/37 pruebas aprobadas"],
        ["Flutter analyze", "0 errores (15 advertencias y 99 informativos, sin impacto en las pruebas)"],
        ["E2E vía API", "38/38 escenarios aprobados"],
        ["CRUD administrativo", "21/21 escenarios aprobados"],
        ["Modelo de IA activo", "lichen_model_v8.keras"],
        ["Inferencias IA", "9 inferencias reales correctas/operativas + 2 errores controlados"],
        ["Seguridad", "BUG-001 corregido y revalidado (Módulo 2)"],
    ]
    table_(doc, ["Resultado", "Valor real"], res, widths=[2.6, 4.2])
    doc.add_paragraph()
    heading(doc, "8.1 Pruebas automatizadas específicas del módulo", 2)
    doc.add_paragraph("Backend (pytest): " + cfg["pytest"]["texto"])
    doc.add_paragraph("Frontend (Flutter): " + cfg["flutter"])
    doc.add_paragraph("E2E vía API: " + cfg["e2e"])
    doc.add_paragraph("CRUD administrativo: " + cfg["crud"])

    # ---------- 9. Hallazgos ----------
    heading(doc, "9. Hallazgos y correcciones", 1)
    if cfg["bug"]:
        b = cfg["bug"]
        table_(doc, ["ID Bug", "Módulo", "Descripción", "Severidad", "Prioridad", "Estado"],
               [[b["id"], b["modulo"], b["descripcion"], b["severidad"], b["prioridad"], "Corregido y revalidado"]],
               widths=[0.8, 0.8, 2.6, 0.8, 0.8, 1.4], font_size=7.0)
        doc.add_paragraph()
        doc.add_paragraph("Corrección aplicada: " + b["correccion"])
        doc.add_paragraph("Revalidación: " + b["revalidacion"])
    else:
        doc.add_paragraph(
            "Durante la ejecución de las pruebas del Sprint 5 no se registraron hallazgos "
            "específicos de este módulo. El defecto de seguridad BUG-001 (control de acceso de "
            "imágenes) se documentó y corrigió en el Módulo 2 (Captura y análisis de imágenes)."
        )

    # ---------- 10. Evidencias ----------
    heading(doc, "10. Evidencias del módulo", 1)
    doc.add_paragraph(
        "Las evidencias correspondientes a este módulo se encuentran en docs/sprint5/evidencias/ "
        "y se enumeran a continuación (archivos reales):"
    )
    for ev in cfg["evidencias"]:
        doc.add_paragraph("- " + ev, style="List Bullet")
    if cfg.get("imagen_ia"):
        fig = EVID / "ia" / "imagen_no_liquen.png"
        if fig.exists():
            doc.add_paragraph()
            figure(doc, fig, "Figura 1. Imagen sintética de prueba (no líquen) utilizada como entrada real en la inferencia del Módulo 2.")

    # ---------- 11. Responsables y cronograma ----------
    heading(doc, "11. Responsables y cronograma", 1)
    doc.add_paragraph(
        "Integrantes del equipo del proyecto (según README): Hugo Andres Mancera Perez, Daniel "
        "Camilo Luque Briceño, Saira Yineth Aragon Suarez, Neyireth Dayana Soriano Ruiz y Heidy "
        "Lizeth Vivas Ramirez. La ejecución y elaboración del Sprint 5 estuvo a cargo del "
        "responsable de QA."
    )
    crono = [
        ["Análisis y auditoría inicial del proyecto", "QA", "01/09/2026"],
        ["Plan y diseño de casos de prueba", "QA", "02/09/2026"],
        ["Ejecución de pruebas manuales y de API", "QA", "02–03/09/2026"],
        ["Pruebas del modelo de IA", "QA", "03/09/2026"],
        ["Pruebas automatizadas (pytest / Flutter)", "QA", "03–05/09/2026"],
        ["Correcciones y regresión", "QA", "04–05/09/2026"],
        ["Elaboración del documento final", "QA", "05/09/2026"],
    ]
    table_(doc, ["Actividad", "Responsable", "Fecha"], crono, widths=[4.4, 1.4, 1.4])

    # ---------- 12. Conclusiones ----------
    heading(doc, "12. Conclusiones del módulo", 1)
    for c in cfg["conclusiones"]:
        doc.add_paragraph(c, style="List Bullet")

    # ---------- 13. Integración ----------
    heading(doc, "13. Integración con el sistema completo", 1)
    doc.add_paragraph(cfg["integracion"])

    # ---------- 14. Consideraciones de plataforma ----------
    heading(doc, "14. Consideraciones de plataforma y estrategia de pruebas", 1)
    doc.add_paragraph(
        "Lichen Dreams se desarrolla y valida como aplicación móvil para Android. Aunque Flutter "
        "es un framework multiplataforma, el soporte de otras plataformas no forma parte del "
        "alcance funcional ni de validación de este proyecto. Las funcionalidades principales "
        "(cámara, permisos, GPS/ubicación, Google Maps móvil, selector de imágenes y ciclo de vida "
        "de la aplicación) dependen de APIs nativas del dispositivo Android."
    )
    doc.add_paragraph(
        "Durante la revisión de herramientas de automatización se evaluó la posibilidad de utilizar "
        "Selenium. Sin embargo, debido a que Lichen Dreams se desarrolla y valida como aplicación "
        "móvil Android, la automatización web no representa el entorno principal de ejecución del "
        "producto. Por esta razón, Selenium no se incorporó como herramienta de validación oficial "
        "del Sprint 5. La automatización y validación futura de pruebas end-to-end móviles puede "
        "realizarse mediante herramientas orientadas al ecosistema móvil, como Flutter "
        "integration_test, Patrol, Appium o Maestro."
    )

    out = OUT_DIR / cfg["archivo"]
    doc.save(str(out))
    return out


if __name__ == "__main__":
    for k in TECNICOS:
        p = build_module_doc(k)
        print("GENERADO:", p)