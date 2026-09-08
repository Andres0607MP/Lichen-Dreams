import sys, zipfile, re
from docx import Document
from pathlib import Path

DOCX = Path(sys.argv[1])
zf = zipfile.ZipFile(str(DOCX))
print("testzip:", "OK" if zf.testzip() is None else "CORRUPTO")
imgs = [n for n in zf.namelist() if n.startswith("word/media/")]
print("imagenes incrustadas:", len(imgs))
for i in imgs:
    print("  -", i)
doc = Document(str(DOCX))
print("parrafos:", len(doc.paragraphs), "| tablas:", len(doc.tables))
txt = "\n".join(p.text for p in doc.paragraphs)
for kw in ["6 pruebas", "interrumpidas", "pruebas web automatizadas", "pruebas de interfaz web automatizada"]:
    print(f'referencia "{kw}":', len(re.findall(kw, txt, re.IGNORECASE)))
print("--- menciones Selenium ---")
for m in re.finditer(r"[^.\n]*Selenium[^.\n]*", txt):
    s = m.group(0).strip()
    print(" *", s[:160])