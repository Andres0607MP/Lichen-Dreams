# -*- coding: utf-8 -*-
"""Validación integral de los 5 DOCX de entregables del Sprint 5."""
import json, re, zipfile
from pathlib import Path
from docx import Document

OUT = Path("docs/sprint5")
FILES = [
    "Sprint_5_Entregable_1_Plan_de_Pruebas.docx",
    "Sprint_5_Entregable_2_Diseno_de_Casos_de_Prueba.docx",
    "Sprint_5_Entregable_3_Matriz_Ejecucion_Pruebas_Manuales.docx",
    "Sprint_5_Entregable_4_Reporte_Hallazgos_Bug_Report.docx",
    "Sprint_5_Entregable_5_Evidencias_Pruebas_Automatizadas.docx",
]
PERMITIDO_CONTEXTO = {"selenium", "automatización web"}
BANNED = ["móvil/web", "web/móvil", "aplicación web/móvil", "aplicación multiplataforma validada",
          "pruebas web automatizadas", "interrumpidas", "Selenium aprobado", "Selenium fallido",
          "Selenium bloqueado", "6 pruebas"]
REQUIRED = ["Aplicación móvil (Android)", "Flutter es un framework multiplataforma", "FastAPI",
            "MySQL", "Android", "176/176", "37/37", "38/38", "21/21", "lichen_model_v8.keras",
            "9 inferencias", "BUG-001", "integration_test", "Patrol", "Appium", "Maestro"]

report = {}
for f in FILES:
    path = OUT / f
    e = {"file": f, "exists": path.exists()}
    if not path.exists():
        report[f] = e
        continue
    zf = zipfile.ZipFile(str(path))
    e["testzip"] = "OK" if zf.testzip() is None else f"CORRUPTO:{zf.testzip()}"
    names = zf.namelist()
    e["document_xml"] = "word/document.xml" in names
    e["styles_xml"] = "word/styles.xml" in names
    e["media"] = [n for n in names if n.startswith("word/media/")]
    doc = Document(str(path))
    blocks = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                blocks.append(c.text)
    full = "\n".join(blocks)
    full = full.replace("\ufffd", "?")
    e["parrafos"] = len(doc.paragraphs)
    e["tablas"] = len(doc.tables)
    e["header"] = doc.sections[0].header.paragraphs[0].text[:90].replace("\ufffd", "?")
    e["pie_num"] = "Página" in (doc.sections[0].footer.paragraphs[0].text or "")
    e["banned_hits"] = [x for x in BANNED if x.lower() in full.lower()]
    e["required_missing"] = [x for x in REQUIRED if x not in full]
    e["selenium_count"] = len(re.findall(r"Selenium", full))
    e["headings"] = len([p for p in doc.paragraphs if p.style.name.startswith("Heading")])
    report[f] = e

Path("docs/sprint5/evidencias/validacion_entregables_docx.json").write_text(
    json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))