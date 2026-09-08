# -*- coding: utf-8 -*-
"""Validación integral de los 5 DOCX del Sprint 5 (macros por módulo)."""
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document

OUT_DIR = Path("docs/sprint5")
FILES = [
    "Sprint_5_Modulo_1_Autenticacion_Usuario.docx",
    "Sprint_5_Modulo_2_Captura_Analisis_IA.docx",
    "Sprint_5_Modulo_3_Resultados_Historial_Reportes.docx",
    "Sprint_5_Modulo_4_Ubicacion_Mapa.docx",
    "Sprint_5_Modulo_5_LiquenPedia_Comunidad_Complementarios.docx",
]

FORBIDDEN = [
    "móvil/web", "web/móvil", "aplicación web", "producto multiplataforma validado",
    "pruebas web", "automatización web como parte de las pruebas oficiales",
]
REQUIRED = [
    "Aplicación móvil (Android)",
    "Flutter es un framework multiplataforma",
    "Android",
    "176/176",
    "37/37",
    "38/38",
    "21/21",
    "lichen_model_v8.keras",
    "9 inferencias reales",
    "BUG-001",
]

report = {}
for f in FILES:
    path = OUT_DIR / f
    entry = {"file": f, "exists": path.exists()}
    if not path.exists():
        report[f] = entry
        continue
    # integridad
    zf = zipfile.ZipFile(str(path))
    entry["testzip"] = "OK" if zf.testzip() is None else zf.testzip()
    names = zf.namelist()
    entry["document_xml"] = "word/document.xml" in names
    entry["styles_xml"] = "word/styles.xml" in names
    entry["media"] = [n for n in names if n.startswith("word/media/")]
    # estructura
    doc = Document(str(path))
    paras = doc.paragraphs
    entry["parrafos"] = len(paras)
    entry["tablas"] = len(doc.tables)
    entry["headings"] = [p.text for p in paras if p.style.name.startswith("Heading")]
    entry["header"] = doc.sections[0].header.paragraphs[0].text[:80]
    entry["footer_pie"] = "Página" in (doc.sections[0].footer.paragraphs[0].text or "")
    # texto completo (párrafos + tablas)
    blocks = [p.text for p in paras]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                blocks.append(c.text)
    full = "\n".join(blocks)
    entry["forbidden_hits"] = [x for x in FORBIDDEN if x.lower() in full.lower()]
    entry["missing_required"] = [x for x in REQUIRED if x not in full]
    entry["selenium_mentions"] = len(re.findall(r"Selenium", full))
    # conteo de páginas aproximado con Word COM si es posible
    report[f] = entry

# Word COM para conteo de páginas (opcional, requiere Word instalado)
try:
    import win32com.client  # pywin32
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    for f in FILES:
        path = str((OUT_DIR / f).resolve())
        d = word.Documents.Open(path, ReadOnly=True)
        report[f]["paginas_word"] = int(d.ComputeStatistics(2))  # wdStatisticPages
        d.Close(False)
    word.Quit()
except Exception as e:
    for f in FILES:
        if f in report:
            report[f]["paginas_word"] = f"no disponible: {e}"

print(json.dumps(report, ensure_ascii=False, indent=2))