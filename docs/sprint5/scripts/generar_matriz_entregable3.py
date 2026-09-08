# -*- coding: utf-8 -*-
"""Genera el ENTREGABLE 3 — Matriz de Ejecución de Pruebas Manuales (Sprint 5).

Productos:
  - docs/sprint5/Sprint_5_Entregable_3_Matriz_Ejecucion_Pruebas_Manuales.xlsx
    (hojas: Resumen, Matriz de Ejecución, Evidencias, Hallazgos)
  - docs/sprint5/Sprint_5_Entregable_3_Matriz_Ejecucion_Pruebas_Manuales.docx

Base: casos del Entregable 2 (TECNICOS) y evidencias reales del proyecto.
Estados de la columna Estado: Aprobado / Fallido / Bloqueado.
"""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from generar_modulos_docx import TECNICOS

ROOT = Path(__file__).resolve().parents[3]
BASE = ROOT / "docs" / "sprint5"
XLSX = BASE / "Sprint_5_Entregable_3_Matriz_Ejecucion_Pruebas_Manuales.xlsx"
DOCX_PATH = BASE / "Sprint_5_Entregable_3_Matriz_Ejecucion_Pruebas_Manuales.docx"

FECHA = "05/09/2026"
GUIA = "GUIA_EVIDENCIAS_MANUALES.txt"
MODSHORT = {"autenticacion": "M1", "captura_ia": "M2", "resultados_historial": "M3",
            "ubicacion_mapa": "M4", "liquenpedia_comunidad": "M5"}
MODNOM = {"autenticacion": "Módulo 1 — Autenticación y gestión de usuario",
          "captura_ia": "Módulo 2 — Captura y análisis de imágenes con IA",
          "resultados_historial": "Módulo 3 — Resultados, historial y reportes",
          "ubicacion_mapa": "Módulo 4 — Ubicación y mapa",
          "liquenpedia_comunidad": "Módulo 5 — LiquenPedia, comunidad y complementarios"}

# Casos cuya evidencia esperada es un pantallazo de la aplicación Android
VISUAL = {
    "M1-01", "M1-02", "M1-03", "M1-04", "M1-05", "M1-06", "M1-07", "M1-08", "M1-09",
    "M1-10", "M1-12", "M1-13", "M1-14", "M1-15", "M1-16", "M1-17", "M1-18", "M1-21",
    "M2-01", "M2-02", "M2-03", "M2-04", "M2-05", "M2-06", "M2-07", "M2-08", "M2-09",
    "M2-13", "M2-14",
    "M3-01", "M3-02", "M3-03", "M3-04", "M3-05", "M3-06", "M3-07", "M3-08", "M3-09",
    "M3-10",
    "M4-01", "M4-02", "M4-03", "M4-04", "M4-08", "M4-09",
    "M5-01", "M5-02", "M5-03", "M5-04", "M5-06", "M5-08", "M5-09", "M5-10", "M5-11",
}


def merge_by_id(module_key):
    """Devuelve filas [ID, Mod, Caso, Pre, Entrada, Esperado, Obtenido, Estado, Evidencia, Obs]."""
    casos = TECNICOS[module_key]["casos"]
    matriz = TECNICOS[module_key]["matriz"]
    cmap = {c[0]: c for c in casos}
    rows = []
    for m in matriz:
        cid = m[0]
        c = cmap.get(cid, None)
        caso_nombre = c[1] if c else m[1]
        pre = c[2] if c else "—"
        entrada = c[3] if c else "—"
        esperado = c[5] if c else m[2]
        obtenido = m[3]
        estado_raw = m[4]
        if estado_raw.startswith("No ejecutado"):
            estado = "Bloqueado"
            obtenido = "No ejecutado"
            obs = m[5] if m[5] else "Pendiente de ejecución manual"
            evidencia = "Pendiente de ejecución manual"
        else:
            estado = estado_raw  # Aprobado
            ev = m[5]
            evidencia = ("Evidencia técnica: " + ev) if ev and not ev.lower().startswith("evidencia técnica") else (ev or "—")
            if cid in VISUAL:
                evidencia += " · Captura manual pendiente (ver " + GUIA + ")"
                obs = "Validado mediante evidencia técnica; captura manual de la aplicación pendiente."
            else:
                obs = "Validado mediante evidencia técnica."
            if not ev:
                evidencia = "Pendiente de evidencia técnica"
                obs = "Revisar evidencia técnica pendiente."
        rows.append([cid, MODSHORT[module_key], caso_nombre, pre, entrada, esperado,
                     obtenido, estado, evidencia, obs])
    return rows


def build_xlsx(all_rows):
    wb = Workbook()

    # ---------- Hoja Resumen ----------
    ws = wb.active
    ws.title = "Resumen"
    info = [
        ["Proyecto", "Lichen Dreams"],
        ["Sprint", "Sprint 5 — Pruebas y Estabilización del Software"],
        ["Entregable", "3 — Matriz de Ejecución de Pruebas Manuales"],
        ["Tecnología", "Flutter"],
        ["Plataforma objetivo", "Android"],
        ["Fecha", FECHA],
    ]
    for i, (k, v) in enumerate(info, start=1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    r0 = len(info) + 2
    n = len(all_rows)
    last = n + 1  # data rows: 2..(n+1)
    mod_range = f"'Matriz de Ejecución'!$B$2:$B${last}"
    est_range = f"'Matriz de Ejecución'!$H$2:$H${last}"
    ws.cell(row=r0, column=1, value="Total de casos").font = Font(bold=True)
    ws.cell(row=r0, column=2, value=f"=COUNTA('Matriz de Ejecución'!$A$2:$A${last})")
    ws.cell(row=r0+1, column=1, value="Casos ejecutados").font = Font(bold=True)
    ws.cell(row=r0+1, column=2, value=f"=COUNTIF({est_range},\"Aprobado\")+COUNTIF({est_range},\"Fallido\")")
    ws.cell(row=r0+2, column=1, value="Aprobados").font = Font(bold=True)
    ws.cell(row=r0+2, column=2, value=f"=COUNTIF({est_range},\"Aprobado\")")
    ws.cell(row=r0+3, column=1, value="Fallidos").font = Font(bold=True)
    ws.cell(row=r0+3, column=2, value=f"=COUNTIF({est_range},\"Fallido\")")
    ws.cell(row=r0+4, column=1, value="Bloqueados / No ejecutados").font = Font(bold=True)
    ws.cell(row=r0+4, column=2, value=f"=COUNTIF({est_range},\"Bloqueado\")")
    ws.cell(row=r0+5, column=1, value="Porcentaje de ejecución").font = Font(bold=True)
    ws.cell(row=r0+5, column=2, value=f"=IFERROR((COUNTIF({est_range},\"Aprobado\")+COUNTIF({est_range},\"Fallido\"))/COUNTA('Matriz de Ejecución'!$A$2:$A${last}),0)")
    ws.cell(row=r0+6, column=1, value="Porcentaje de aprobación").font = Font(bold=True)
    ws.cell(row=r0+6, column=2, value=f"=IFERROR(COUNTIF({est_range},\"Aprobado\")/(COUNTIF({est_range},\"Aprobado\")+COUNTIF({est_range},\"Fallido\")),0)")
    for c in [1, 2]:
        ws.cell(row=r0+5, column=c).number_format = "0.0%"
        ws.cell(row=r0+6, column=c).number_format = "0.0%"

    hdr_res = r0 + 8
    ws.cell(row=hdr_res, column=1, value="Módulo").font = Font(bold=True)
    ws.cell(row=hdr_res, column=2, value="Total").font = Font(bold=True)
    ws.cell(row=hdr_res, column=3, value="Ejecutados").font = Font(bold=True)
    ws.cell(row=hdr_res, column=4, value="Aprobados").font = Font(bold=True)
    ws.cell(row=hdr_res, column=5, value="Fallidos").font = Font(bold=True)
    ws.cell(row=hdr_res, column=6, value="Bloqueados").font = Font(bold=True)
    for i, mod in enumerate(["M1", "M2", "M3", "M4", "M5"], start=1):
        r = hdr_res + i
        ws.cell(row=r, column=1, value=mod)
        ws.cell(row=r, column=2, value=f"=COUNTIFS({mod_range},$A{r},{est_range},\"*\")")
        ws.cell(row=r, column=3, value=f"=COUNTIFS({mod_range},$A{r},{est_range},\"Aprobado\")+COUNTIFS({mod_range},$A{r},{est_range},\"Fallido\")")
        ws.cell(row=r, column=4, value=f"=COUNTIFS({mod_range},$A{r},{est_range},\"Aprobado\")")
        ws.cell(row=r, column=5, value=f"=COUNTIFS({mod_range},$A{r},{est_range},\"Fallido\")")
        ws.cell(row=r, column=6, value=f"=COUNTIFS({mod_range},$A{r},{est_range},\"Bloqueado\")")
    for col, w in zip("ABCDEF", [14, 16, 16, 16, 16, 24]):
        ws.column_dimensions[col].width = w

    # ---------- Hoja Matriz de Ejecución ----------
    ws2 = wb.create_sheet("Matriz de Ejecución")
    headers = ["ID", "Módulo", "Caso de prueba", "Precondición", "Entrada",
               "Resultado esperado", "Resultado obtenido", "Estado", "Evidencia", "Observaciones"]
    hf = PatternFill("solid", fgColor="1F4E79")
    for j, h in enumerate(headers, start=1):
        c = ws2.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hf
    for i, row in enumerate(all_rows, start=2):
        for j, val in enumerate(row, start=1):
            ws2.cell(row=i, column=j, value=val)
    widths = [8, 8, 26, 22, 24, 26, 20, 12, 42, 34]
    for j, w in enumerate(widths, start=1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    for i in range(2, last + 1):
        st = ws2.cell(row=i, column=8).value
        fill = None
        if st == "Aprobado":
            fill = PatternFill("solid", fgColor="E2EFDA")
        elif st == "Fallido":
            fill = PatternFill("solid", fgColor="FCE4EC")
        elif st == "Bloqueado":
            fill = PatternFill("solid", fgColor="FFF2CC")
        if fill:
            ws2.cell(row=i, column=8).fill = fill
        ws2.cell(row=i, column=1).alignment = Alignment(horizontal="center")
    ws2.freeze_panes = "A2"

    # ---------- Hoja Evidencias ----------
    ws3 = wb.create_sheet("Evidencias")
    for j, h in enumerate(["ID Caso", "Tipo de evidencia", "Archivo / referencia", "Descripción"], start=1):
        c = ws3.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hf
    erow = 2
    for r in all_rows:
        cid, _, caso, _, _, _, _, estado, ev, obs = r
        if estado == "Bloqueado":
            ws3.cell(row=erow, column=1, value=cid)
            ws3.cell(row=erow, column=2, value="Captura de pantalla / ejecución manual")
            ws3.cell(row=erow, column=3, value="Pendiente de ejecución manual")
            ws3.cell(row=erow, column=4, value=f"{caso} — {obs}")
            erow += 1
            continue
        if "frontend/test" in ev:
            tipo = "Widget test"
        elif ev.startswith("Evidencia técnica: ia/"):
            tipo = "Modelo IA"
        elif "correcciones/" in ev:
            tipo = "Backend"
        elif "backend/tests" in ev:
            tipo = "Backend"
        elif "api/" in ev:
            tipo = "API / JSON"
        else:
            tipo = "Otra"
        ref = ev.replace("Evidencia técnica: ", "").split(" · ")[0] if ev else "—"
        ws3.cell(row=erow, column=1, value=cid)
        ws3.cell(row=erow, column=2, value=tipo)
        ws3.cell(row=erow, column=3, value=ref if ref != "Pendiente de ejecución manual" else "Pendiente")
        ws3.cell(row=erow, column=4, value=f"{caso} — {obs}")
        erow += 1
        if cid in VISUAL:
            ws3.cell(row=erow, column=1, value=cid)
            ws3.cell(row=erow, column=2, value="Captura de pantalla")
            ws3.cell(row=erow, column=3, value="Pendiente de captura manual")
            ws3.cell(row=erow, column=4, value="Pantallazo de la aplicación Android a tomar según " + GUIA)
            erow += 1
    for col, w in zip("ABCD", [10, 30, 40, 60]):
        ws3.column_dimensions[col].width = w
    ws3.freeze_panes = "A2"

    # ---------- Hoja Hallazgos ----------
    ws4 = wb.create_sheet("Hallazgos")
    for j, h in enumerate(["ID Bug", "Caso asociado", "Módulo", "Descripción", "Severidad", "Prioridad", "Estado"], start=1):
        c = ws4.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hf
    ws4.append(["BUG-001", "M2-11 (y servicio de imágenes)", "Módulo 2 — Imágenes",
                "Control de acceso: los endpoints /imagenes, /imagenes/{id} y DELETE /imagenes/{id} "
                "no exigían autenticación y exponían rutas de imágenes privadas (analyses/user_N), "
                "permitiendo eliminarlas sin sesión. Corregido protegiendo los tres endpoints con "
                "autenticación y control de propiedad (CAN_VIEW_PRIVATE_IMAGES) y revalidado "
                "(401/200/204/403 y prueba de regresión).",
                "High", "High", "Corrected"])
    for col, w in zip("ABCDEFG", [12, 22, 24, 80, 12, 12, 14]):
        ws4.column_dimensions[col].width = w
    ws4.freeze_panes = "A2"

    wb.save(str(XLSX))


# ---------------------------------------------------------------------------
# DOCX Entregable 3
# ---------------------------------------------------------------------------
AZUL = RGBColor(0x1F, 0x4E, 0x79)
GRIS = RGBColor(0x59, 0x59, 0x59)


def set_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def hf(doc):
    s = doc.sections[0]
    hp = s.header.paragraphs[0]
    hp.text = "Lichen Dreams — Sprint 5 · Matriz de Ejecución de Pruebas Manuales"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.size = Pt(8); r.font.color.rgb = GRIS; r.font.italic = True
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Página ")
    run.font.size = Pt(8); run.font.color.rgb = GRIS
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r1 = fp.add_run(); r1.font.size = Pt(8); r1._r.append(f1)
    r2 = fp.add_run(); r2._r.append(it)
    r3 = fp.add_run(); r3.font.size = Pt(8); r3._r.append(f2)
    for rr in [r1, r2, r3]:
        rr.font.color.rgb = GRIS


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = AZUL
    return h


def table_(doc, headers, rows, widths=None, font_size=7.0):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htext in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(htext)
        r.bold = True; r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_bg(cell, "1F4E79")
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(str(val))
            r.font.size = Pt(font_size)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Inches(w)


def portada(doc):
    doc.add_paragraph(); doc.add_paragraph()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("LICHEN DREAMS"); r.font.size = Pt(30); r.bold = True; r.font.color.rgb = AZUL
    st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = st.add_run("Sprint 5 — Pruebas y Estabilización del Software"); sr.font.size = Pt(18); sr.bold = True
    eg = doc.add_paragraph(); eg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    er = eg.add_run("ENTREGABLE 3 — MATRIZ DE EJECUCIÓN DE PRUEBAS MANUALES")
    er.font.size = Pt(15); er.font.color.rgb = GRIS
    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        "Tecnología principal: Flutter · Plataforma objetivo: Android\n\n"
        "Aprendices SENA · Análisis y Desarrollo de Software (ADSO)\n\nFecha: 05/09/2026"
    )
    mr.font.size = Pt(10); mr.font.color.rgb = GRIS
    doc.add_page_break()


def build_docx(all_rows):
    doc = Document(); hf(doc)
    portada(doc)

    heading(doc, "1. Información general", 1)
    table_(doc, ["Campo", "Información"], [
        ["Proyecto", "Lichen Dreams"],
        ["Sprint", "Sprint 5 — Pruebas y Estabilización del Software"],
        ["Entregable", "3 — Matriz de Ejecución de Pruebas Manuales"],
        ["Tipo de sistema", "Aplicación móvil (Android) desarrollada con Flutter"],
        ["Tecnología", "Flutter · Backend FastAPI · Base de datos MySQL"],
        ["Plataforma objetivo", "Android"],
        ["Fecha", FECHA],
    ], widths=[2.2, 4.6])

    heading(doc, "2. Objetivo", 1)
    doc.add_paragraph(
        "Registrar el resultado real de la ejecución de cada caso de prueba definido en el "
        "Entregable 2, indicando el estado (Aprobado, Fallido o Bloqueado), la evidencia real que "
        "lo respalda y las observaciones correspondientes, de modo que la matriz sea verificable "
        "y coherente con la ejecución efectiva del Sprint 5."
    )
    heading(doc, "3. Criterios de ejecución", 1)
    doc.add_paragraph(
        "- Aprobado: el caso fue ejecutado y el comportamiento obtenido coincide con el esperado.\n"
        "- Fallido: el caso fue ejecutado y se observó un comportamiento incorrecto.\n"
        "- Bloqueado: el caso no pudo ejecutarse por una condición externa real (dispositivo, "
        "configuración, dependencia no disponible); en la matriz se indica el motivo.\n"
        "- Evidencia técnica vs. visual: cuando la validación fue por API/backend/automatizada se "
        "indica «Evidencia técnica: …». Las capturas de la aplicación Android se tomarán según la "
        "guía " + GUIA + " y se guardarán en docs/sprint5/evidencias/manuales/.\n"
        "- No se marca un caso como aprobado por el solo hecho de que la funcionalidad exista: se "
        "requiere ejecución y evidencia real."
    )
    doc.add_paragraph(
        "Consideración sobre automatización web: Selenium se evaluó de forma exploratoria, pero no "
        "se adopta como validación oficial porque el producto se desarrolla y valida como "
        "aplicación móvil Android (las funcionalidades de cámara, GPS, mapa, permisos y "
        "autenticación móvil no se ejercen en un build web)."
    )

    heading(doc, "4. Matriz de ejecución de pruebas manuales", 1)
    doc.add_paragraph("Casos M1 a M5 del Entregable 2 con su resultado de ejecución real.")
    from collections import Counter
    for key, short in MODSHORT.items():
        heading(doc, f"4.{int(short[1])} {MODNOM[key]}", 2)
        filas = []
        for r in all_rows:
            if r[1] != short:
                continue
            filas.append([r[0], r[2], r[5], r[6], r[7], r[8]])
        table_(doc, ["ID", "Caso de prueba", "Resultado esperado", "Resultado obtenido", "Estado", "Evidencia"],
               filas, widths=[0.6, 1.6, 1.5, 1.2, 0.8, 2.1], font_size=6.0)
        doc.add_paragraph()

    heading(doc, "5. Resumen de resultados", 1)
    stats = []
    total = len(all_rows)
    for key, short in MODSHORT.items():
        mod_rows = [r for r in all_rows if r[1] == short]
        ap = sum(1 for r in mod_rows if r[7] == "Aprobado")
        fa = sum(1 for r in mod_rows if r[7] == "Fallido")
        bl = sum(1 for r in mod_rows if r[7] == "Bloqueado")
        stats.append([MODNOM[key], len(mod_rows), ap + fa, ap, fa, bl])
    table_(doc, ["Módulo", "Total", "Ejecutados", "Aprobados", "Fallidos", "Bloqueados"],
           stats, widths=[2.8, 0.8, 0.9, 0.9, 0.8, 0.9], font_size=7.5)
    tot_ap = sum(s[3] for s in stats)
    tot_fa = sum(s[4] for s in stats)
    tot_bl = sum(s[5] for s in stats)
    doc.add_paragraph()
    doc.add_paragraph(
        f"Totales: {total} casos · {tot_ap + tot_fa} ejecutados · {tot_ap} aprobados · "
        f"{tot_fa} fallidos · {tot_bl} bloqueados. Porcentaje de ejecución: "
        f"{100 * (tot_ap + tot_fa) / total:.1f}%. Porcentaje de aprobación (sobre ejecutados): "
        f"{100 * tot_ap / max(tot_ap + tot_fa, 1):.1f}%."
    )
    doc.add_paragraph(
        "Resultados oficiales complementarios del Sprint 5 (conservados): backend pytest 176/176; "
        "Flutter 37/37; flutter analyze 0 errores; E2E API 38/38; CRUD administrativo 21/21; "
        "modelo IA activo lichen_model_v8.keras con 9 inferencias reales + 2 errores controlados; "
        "BUG-001 corregido y revalidado."
    )

    heading(doc, "6. Evidencias", 1)
    doc.add_paragraph(
        "Evidencias técnicas reales referenciadas en la matriz (carpeta docs/sprint5/evidencias/):"
    )
    for f in [
        "api/auditoria_sprint5_dllpsh.json (auditoría E2E vía API, 38 escenarios aprobados)",
        "api/auditoria_crud_admin.json y api/ejecucion_crud_admin.txt (auditoría CRUD administrativa, 21 escenarios)",
        "ia/resultados_ia.json y ia/ejecucion_ia.txt (inferencias del modelo v8)",
        "ia/imagen_no_liquen.png (imagen de prueba no-líquen utilizada en la inferencia)",
        "correcciones/resultado_correccion_bug001.json y correcciones/ejecucion_bug001.txt (validación BUG-001)",
        "backend/tests/ (suite pytest 176/176) y frontend/test/ (suite Flutter 37/37)",
    ]:
        doc.add_paragraph("- " + f, style="List Bullet")
    doc.add_paragraph(
        "Las capturas de pantalla de la aplicación Android se almacenarán en "
        "docs/sprint5/evidencias/manuales/ siguiendo la guía " + GUIA + ". En la matriz, los casos "
        "visuales figuran como «Captura manual pendiente» hasta que exista el pantallazo real; no "
        "se inventan nombres de archivos que no existen todavía."
    )

    heading(doc, "7. Hallazgos", 1)
    table_(doc, ["ID Bug", "Caso Asociado", "Módulo", "Descripción", "Severidad", "Prioridad", "Estado"], [
        ["BUG-001", "M2-11 (servicio de imágenes)", "Módulo 2 — Imágenes",
         "Control de acceso de imágenes: endpoints /imagenes, /imagenes/{id} y DELETE /imagenes/{id} "
         "sin autenticación, exponiendo rutas privadas y permitiendo eliminación ajena. Corregido "
         "con autenticación y control de propiedad; revalidado (401/200/204/403) con prueba de regresión.",
         "High", "High", "Corrected"],
    ], widths=[0.7, 1.1, 1.2, 2.6, 0.7, 0.7, 0.8], font_size=6.5)
    doc.add_paragraph(
        "No se registraron otros defectos reales durante el Sprint 5."
    )

    heading(doc, "8. Conclusiones", 1)
    doc.add_paragraph(
        "Los casos con ejecución real (vía API, backend, pruebas automatizadas e IA) quedaron "
        "registrados como Aprobados con su evidencia técnica. Los casos que requieren la aplicación "
        "Android en dispositivo (cámara, galería, GPS, Google Maps, Google Sign-In en dispositivo, "
        "UI administrativa) quedaron como Bloqueados, pendientes de ejecución manual y captura, "
        "sin resultados inventados. La matriz será complementada con los pantallazos reales según "
        "la guía, manteniendo la coherencia entre la aplicación, el backend y la documentación."
    )

    return doc


def main():
    global DOCX_PATH
    if len(sys.argv) > 1:
        DOCX_PATH = Path(sys.argv[1]) / DOCX_PATH.name
    all_rows = []
    for key in TECNICOS:
        all_rows.extend(merge_by_id(key))
    build_xlsx(all_rows)
    doc = build_docx(all_rows)
    doc.save(str(DOCX_PATH))
    print("XLSX:", XLSX)
    print("DOCX:", DOCX_PATH)
    print("total casos:", len(all_rows))


if __name__ == "__main__":
    main()