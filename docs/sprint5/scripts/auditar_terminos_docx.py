# -*- coding: utf-8 -*-
"""Auditoría terminológica del documento Sprint 5."""
import re
import sys
from docx import Document


def load_text(path):
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(("P", p.style.name, p.text))
    for ti, t in enumerate(doc.tables):
        for r in t.rows:
            for c in r.cells:
                if c.text.strip():
                    parts.append(("T", f"tabla{ti+1}", c.text))
    return parts


def main(path):
    parts = load_text(path)
    print(f"=== AUDITORÍA TERMINOLÓGICA: {path} ===")
    terms = ["multiplataforma", "plataforma", "móvil", "android", "web", "flutter web", "selenium",
             "pruebas web", "automatización web", "windows", "linux", "macos", "ios", "escritorio",
             "dispositivo", "cámara", "gps", "ubicación"]
    for kind, loc, text in parts:
        low = text.lower()
        found = [t for t in terms if t in low]
        if found:
            print(f"\n[{kind}|{loc}] TERMS={found}")
            print("   ", text[:260].replace("\n", " "))


if __name__ == "__main__":
    main(sys.argv[1])