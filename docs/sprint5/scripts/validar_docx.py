# -*- coding: utf-8 -*-
"""Valida el documento Word generado: integridad, imágenes, tablas y contenido."""
import zipfile
from docx import Document
from pathlib import Path

DOCX = Path("Sprint_5_Pruebas_y_Estabilizacion_Lichen_Dreams.docx")

print("== 1. Integridad ZIP ==")
zf = zipfile.ZipFile(str(DOCX))
bad = zf.testzip()
print("testzip:", "OK (sin archivos corruptos)" if bad is None else f"CORRUPTO: {bad}")

images = [n for n in zf.namelist() if n.startswith("word/media/")]
print("imágenes incrustadas:", len(images))
for im in images:
    print("  -", im)

doc = Document(str(DOCX))
print("\n== 2. Estructura ==")
print("párrafos:", len(doc.paragraphs))
print("tablas:", len(doc.tables))

print("\n== 3. Encabezados (títulos) ==")
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        print(f"  [{p.style.name}] {p.text}")

print("\n== 4. Descripción de tablas ==")
for i, t in enumerate(doc.tables):
    hdr = [c.text for c in t.rows[0].cells]
    print(f"  Tabla {i+1}: {len(t.rows)} filas x {len(t.columns)} cols | encabezados: {hdr}")

print("\n== 5. Referencias a figuras ==")
figs = [p.text for p in doc.paragraphs if p.text.startswith("Figura")]
for f in figs:
    print("  -", f)

print("\n== 6. Pies de página / encabezado ==")
for s in doc.sections:
    print("  header:", s.header.paragraphs[0].text[:80])
    print("  footer:", s.footer.paragraphs[0].text if s.footer.paragraphs[0].text else "(tiene field PAGE)")